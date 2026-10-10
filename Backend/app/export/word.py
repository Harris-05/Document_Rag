"""Word (.docx) export."""

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.export.content import GENERAL_NOTE, VERIFIED_NOTE, ExportContent

_GREEN = RGBColor(0x1B, 0x7A, 0x43)
_RED = RGBColor(0xB4, 0x2B, 0x2B)
_GREY = RGBColor(0x5F, 0x66, 0x70)


def _lines(document, text: str, style: str | None = None) -> None:
    """One paragraph per block, with single line breaks inside it kept as line breaks."""
    paragraph = document.add_paragraph(style=style)
    for index, line in enumerate(text.split("\n")):
        if index:
            paragraph.add_line_break()
        paragraph.add_run(line)


def render_docx(content: ExportContent) -> bytes:
    document = Document()
    document.core_properties.title = "Answer with verified quotes"
    document.core_properties.author = "Marginalia"
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    document.add_heading("Answer with quotes", level=0)
    meta = document.add_paragraph()
    run = meta.add_run(
        f"{'Documents' if len(content.documents) != 1 else 'Document'}: {', '.join(content.documents) or 'unknown'}"
        f"\nExported {content.generated}"
        + (f"\nEvidence: {content.evidence}" if content.evidence else "")
    )
    run.font.size = Pt(9.5)
    run.font.color.rgb = _GREY

    document.add_heading("Question", level=1)
    _lines(document, content.question)

    document.add_heading("Answer", level=1)
    if content.incomplete:
        note = document.add_paragraph()
        flag = note.add_run("This answer was stopped before it finished and may be incomplete.")
        flag.italic = True
        flag.font.color.rgb = _RED
    for paragraph in content.paragraphs or ["(no answer text)"]:
        _lines(document, paragraph)

    if content.explanation:
        document.add_heading("General explanation (not from the document)", level=2)
        note = document.add_paragraph()
        small = note.add_run(GENERAL_NOTE)
        small.italic = True
        small.font.size = Pt(9.5)
        small.font.color.rgb = _GREY
        _lines(document, content.explanation)

    document.add_heading("Quotes", level=1)
    if not content.quotes:
        document.add_paragraph("This answer has no quotes.")
    for quote in content.quotes:
        heading = document.add_paragraph()
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.space_before = Pt(8)
        number = heading.add_run(f"[{quote.n}]  ")
        number.bold = True
        status = heading.add_run("Verified" if quote.verified else "Not verified")
        status.bold = True
        status.font.color.rgb = _GREEN if quote.verified else _RED
        where = []
        if quote.document:
            where.append(quote.document)
        if quote.verified and quote.page:
            where.append(f"page {quote.page}")
        if where:
            place = heading.add_run("   " + ", ".join(where))
            place.font.size = Pt(9.5)
            place.font.color.rgb = _GREY
        body = document.add_paragraph()
        body.paragraph_format.left_indent = Pt(18)
        text = body.add_run(f"“{quote.text}”" if quote.text else "(no quote was given for this citation)")
        text.italic = True

    footer = document.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer.paragraph_format.space_before = Pt(14)
    small = footer.add_run(VERIFIED_NOTE)
    small.font.size = Pt(9)
    small.font.color.rgb = _GREY

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
